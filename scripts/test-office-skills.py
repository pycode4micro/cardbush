"""Behavior tests for the original CardBush XLSX conversion helper."""
import io
import json
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "assets/skills/xlsx/scripts"))
import recalc


def fixture(path, cells='<c r="A1"><f>1+1</f><v/></c>', additions=None):
    main = recalc.SHEET_NS[1:-1]
    relationship = recalc.REL_NS[1:-1]
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as package:
        package.writestr("xl/workbook.xml", f'<workbook xmlns="{main}" xmlns:r="{relationship}"><sheets><sheet name="数据" sheetId="1" r:id="rId1"/></sheets></workbook>')
        package.writestr("xl/_rels/workbook.xml.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Target="worksheets/sheet1.xml"/></Relationships>')
        package.writestr("xl/worksheets/sheet1.xml", f'<worksheet xmlns="{main}"><dimension ref="A1:XFD1048576"/><sheetData><row r="1">{cells}</row></sheetData></worksheet>')
        for name, content in (additions or {}).items():
            package.writestr(name, content)


class WorkbookConversionTests(unittest.TestCase):
    def setUp(self):
        directory = tempfile.TemporaryDirectory(prefix="cardbush-skill-test-")
        self.addCleanup(directory.cleanup)
        self.root = Path(directory.name)
        self.source = self.root / "输入.xlsx"
        self.output = self.root / "输出.xlsx"
        fixture(self.source)
        self.original = self.source.read_bytes()

    def convert(self, runner):
        with patch.object(recalc, "find_soffice", return_value="fixture-engine"), patch.object(recalc.subprocess, "run", side_effect=runner):
            return recalc.recalculate(self.source, self.output, timeout=0.5)

    def write_result(self, args, **kwargs):
        self.assertEqual(kwargs["timeout"], 0.5)
        self.assertNotIn(str(self.source), args, "engine must receive a temporary copy")
        target = Path(args[args.index("--outdir") + 1]) / Path(args[-1]).name
        fixture(target, '<c r="A1"><f>1+1</f><v>2</v></c>')
        return subprocess.CompletedProcess(args, 0, "", "")

    def assert_failure_preserves_files(self, runner):
        with self.assertRaises((RuntimeError, OSError, ValueError, subprocess.SubprocessError)):
            self.convert(runner)
        self.assertEqual(self.source.read_bytes(), self.original)
        self.assertFalse(self.output.exists())
        self.assertFalse(list(self.root.glob("cardbush-xlsx-*")))

    def test_success_creates_new_file_and_preserves_source(self):
        report = self.convert(self.write_result)
        self.assertEqual(report["status"], "success")
        self.assertEqual(report["formula_count"], 1)
        self.assertEqual(recalc.scan_workbook(self.output)["missing_cache_count"], 0)
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_timeout_and_failed_exit_preserve_input(self):
        def timeout(args, **kwargs):
            raise subprocess.TimeoutExpired(args, kwargs["timeout"])
        self.assert_failure_preserves_files(timeout)
        self.assert_failure_preserves_files(lambda args, **kw: subprocess.CompletedProcess(args, 124, "", "failed"))

    def test_exit_zero_without_a_new_workbook_is_not_success(self):
        self.assert_failure_preserves_files(lambda args, **kw: subprocess.CompletedProcess(args, 0, "", ""))

    def test_incomplete_or_erroneous_formula_output_is_not_published(self):
        for cells in ['<c r="A1"><f>1+1</f><v/></c>', '<c r="A1" t="e"><f>1/0</f><v>#DIV/0!</v></c>', '<c r="B1"><f>1+1</f><v>2</v></c>']:
            def runner(args, **kwargs):
                result = self.write_result(args, **kwargs)
                fixture(Path(args[args.index("--outdir") + 1]) / "workbook.xlsx", cells)
                return result
            self.assert_failure_preserves_files(runner)

    def test_existing_output_and_in_place_conversion_are_refused(self):
        self.output.write_bytes(b"keep existing output")
        for destination in [self.output, self.source]:
            with self.assertRaises(ValueError):
                recalc.recalculate(self.source, destination)
        self.assertEqual(self.output.read_bytes(), b"keep existing output")
        self.assertEqual(self.source.read_bytes(), self.original)

    def test_output_created_during_conversion_is_preserved(self):
        def runner(args, **kwargs):
            result = self.write_result(args, **kwargs)
            self.output.write_bytes(b"concurrent edit")
            return result
        with self.assertRaises(FileExistsError):
            self.convert(runner)
        self.assertEqual(self.output.read_bytes(), b"concurrent edit")

    def test_source_changed_during_conversion_is_preserved(self):
        def runner(args, **kwargs):
            result = self.write_result(args, **kwargs)
            self.source.write_bytes(b"new user content")
            return result
        with self.assertRaisesRegex(RuntimeError, "Source changed"):
            self.convert(runner)
        self.assertEqual(self.source.read_bytes(), b"new user content")
        self.assertFalse(self.output.exists())

    def test_empty_string_cache_is_valid_but_error_type_is_detected(self):
        fixture(self.source, '<c r="A1" t="str"><f>""</f><v/></c><c r="B1" t="inlineStr"><is><t>#REF! is documentation</t></is></c>')
        report = recalc.scan_workbook(self.source)
        self.assertEqual(report["formula_count"], 1)
        self.assertEqual(report["missing_cache_count"], 0)
        self.assertEqual(report["error_count"], 0)

    def test_unsafe_round_trips_are_refused(self):
        for member in ["xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml", "xl/connections.xml", "xl/embeddings/oleObject1.bin", "_xmlsignatures/sig1.xml"]:
            fixture(self.source, additions={member: "restricted"})
            with self.assertRaisesRegex(ValueError, "unsupported"):
                recalc.scan_workbook(self.source)
        for timeout in [0, -1, float("nan"), float("inf")]:
            with self.assertRaises(ValueError):
                recalc.recalculate(self.source, self.output, timeout)

    def test_external_service_formulas_are_refused_before_engine_start(self):
        fixture(self.source, '<c r="A1"><f>_xlfn.WEBSERVICE("https://example.test")</f><v/></c>')
        with patch.object(recalc, "find_soffice", side_effect=AssertionError("must not start")):
            with self.assertRaisesRegex(ValueError, "external services"):
                recalc.recalculate(self.source, self.output)

    def test_literal_only_file_does_not_require_calculation_engine(self):
        fixture(self.source, '<c r="A1"><v>42</v></c>')
        with patch.object(recalc, "find_soffice", side_effect=AssertionError("must not start")):
            result = recalc.recalculate(self.source, self.output)
        self.assertEqual(result["status"], "not_needed")
        self.assertEqual(self.source.read_bytes(), self.output.read_bytes())

    def test_dtd_declarations_split_across_chunks_are_rejected(self):
        reader = recalc.XMLReader(io.BytesIO(b'<!DOCTYPE x [<!ENTITY y "value">]>'))
        with self.assertRaises(ValueError):
            while reader.read(2):
                pass

    def test_cli_failure_is_json_and_nonzero(self):
        command = [sys.executable, "-B", str(ROOT / "assets/skills/xlsx/scripts/recalc.py"), str(self.source), str(self.output), "--soffice", str(self.root / "missing-engine")]
        result = subprocess.run(command, capture_output=True, text=True, timeout=10, encoding="utf-8")
        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(json.loads(result.stdout)["status"], "error")
        self.assertFalse(self.output.exists())

    def test_real_libreoffice_when_available(self):
        try:
            executable = recalc.find_soffice()
            from openpyxl import Workbook, load_workbook
        except (FileNotFoundError, ImportError):
            self.skipTest("LibreOffice/openpyxl unavailable; actual calculation not verified")
        book = Workbook()
        book.active["A1"] = "=1+1"
        book.save(self.source)
        book.close()
        recalc.recalculate(self.source, self.output, executable=executable)
        values = load_workbook(self.output, data_only=True)
        self.assertEqual(values.active["A1"].value, 2)
        values.close()


if __name__ == "__main__":
    unittest.main(verbosity=2)
